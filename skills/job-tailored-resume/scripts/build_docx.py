#!/usr/bin/env python3
"""Build a resume or cover letter .docx (and optional PDF) in a clean, ATS-friendly house style.

Usage:
    python build_docx.py spec.json out.docx [--pdf] [--profile profile.json]

The candidate's name and contact line come from profile.json (see templates/profile.example.json).
Lookup order: --profile, $CAREER_AGENCY_HOME/profile.json, ./profile.json, ~/.career-agency/profile.json.
A spec header may override with "name" and/or "contact_items".

The spec is JSON. Inline **bold** markup is supported in every text field.

{
  "header": {
    "name": "optional; defaults to profile.json name",
    "tagline": "optional one-line positioning statement under the name",
    "contact": true            # default true: the contact line from profile.json (email | mobile | links)
  },
  "sections": [
    {"title": "PROFILE", "type": "paragraphs", "items": ["text", "text"]},
    {"title": "CAREER HIGHLIGHTS", "type": "bullets", "items": ["**Lead-in:** detail"]},
    {"title": "PROFESSIONAL EXPERIENCE", "type": "roles", "items": [
        {"title": "Senior Product Manager",
         "org_line": "Northwind Health | Full-time | Seattle, WA, USA | Jun 2021 – Nov 2024",
         "summary": "optional one-line paragraph before bullets",
         "break_before": false,   # true starts this role on a new page (fixes a stranded header)
         "bullets": ["**Owned** ...", "..."]}
    ]},
    {"title": "SKILLS", "type": "labeled", "items": [{"label": "Product", "text": "..."}]},
    {"title": null, "type": "paragraphs", "items": ["letter body paragraph"]}   # cover letter
  ]
}
Any section may also set "break_before": true to start on a new page.

Styling is built in code: single column, no tables or text boxes, 0.5" margins, Arial name and
section headings in dark blue, blue bold role lines, Calibri 11 body, real Word bullets.
"""
import argparse
import json
import os
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BLUE = RGBColor(0x2F, 0x54, 0x96)          # role lines, contact values
HEADING_BLUE = RGBColor(0x1F, 0x4E, 0x79)  # name and section headings

def find_profile(explicit=None):
    """Load profile.json (name + contact items). Returns {} when none is found."""
    if explicit:
        path = Path(explicit).expanduser()
        if not path.is_file():
            sys.exit(f"profile not found: {path}")
        return json.loads(path.read_text(encoding="utf-8"))
    candidates = []
    home = os.environ.get("CAREER_AGENCY_HOME")
    if home:
        candidates.append(Path(home).expanduser() / "profile.json")
    candidates += [Path.cwd() / "profile.json", Path.home() / ".career-agency" / "profile.json"]
    for c in candidates:
        if c.is_file():
            return json.loads(c.read_text(encoding="utf-8"))
    return {}


def contact_items(profile, hdr):
    """Normalise contact entries to (label, text, url) tuples."""
    raw = hdr.get("contact_items") or profile.get("contact") or []
    items = []
    for c in raw:
        if "label" not in c or "text" not in c:
            sys.exit(f"contact entry needs 'label' and 'text': {c}")
        label = c["label"] if c["label"].endswith(":") else c["label"] + ":"
        items.append((label, c["text"], c.get("url")))
    return items


def add_runs(par, text, bold_all=False, color=None, size=None):
    """Add text with **bold** markup as runs."""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = bold_all
        if part.startswith("**") and part.endswith("**"):
            part, bold = part[2:-2], True
        r = par.add_run(part)
        r.bold = bold or None
        if color:
            r.font.color.rgb = color
        if size:
            r.font.size = Pt(size)
    return par


def add_hyperlink(par, url, text, size):
    part = par.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "2F5496"); rpr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); run.append(t)
    h.append(run)
    par._p.append(h)


def set_spacing(par, before=0, after=0):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_spacing(p, 0, 2)
    add_runs(p, text)
    return p


def heading(doc, title):
    h = doc.add_paragraph(style="Heading 2")
    set_spacing(h, 10, 3)
    add_runs(h, title, bold_all=True)
    return h


def new_document():
    """Blank document: 0.5" margins, Calibri 11 body,
    Arial headings in dark blue, real Word bullets."""
    doc = Document()
    for sec in doc.sections:
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(0.5)
    st = doc.styles
    normal = st["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    for name, size in (("Heading 1", 24), ("Heading 2", 13)):
        h = st[name]
        h.font.name = "Arial"; h.font.size = Pt(size); h.font.bold = True
        h.font.color.rgb = HEADING_BLUE
        rfonts = h.element.rPr.rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)
        rfonts.set(qn("w:ascii"), "Arial"); rfonts.set(qn("w:hAnsi"), "Arial")
    return doc


def build(spec, out, profile=None):
    profile = profile or {}
    doc = new_document()
    hdr = spec.get("header", {})
    person = hdr.get("name") or profile.get("name")
    if not person:
        sys.exit("no name: set header.name in the spec or name in profile.json")

    name = doc.add_paragraph(style="Heading 1")
    set_spacing(name, 0, 0)
    r = name.add_run(person)
    r.bold = True; r.font.size = Pt(24); r.font.name = "Arial"

    if hdr.get("tagline"):
        t = doc.add_paragraph()
        set_spacing(t, 0, 2)
        add_runs(t, hdr["tagline"], bold_all=True, color=BLUE, size=10.5)

    items = contact_items(profile, hdr)
    if hdr.get("contact", True) and items:
        c = doc.add_paragraph()
        set_spacing(c, 0, 4)
        for i, (label, text, url) in enumerate(items):
            if i:
                add_runs(c, " | ", size=9)
            add_runs(c, label + " ", bold_all=True, size=9)
            if url:
                add_hyperlink(c, url, text, 9)
            else:
                add_runs(c, text, size=9, color=BLUE)

    for sec in spec.get("sections", []):
        if sec.get("title"):
            h = heading(doc, sec["title"])
            h.paragraph_format.page_break_before = bool(sec.get("break_before"))
        kind = sec.get("type", "paragraphs")
        items = sec.get("items", [])
        if kind == "paragraphs":
            for i, text in enumerate(items):
                p = doc.add_paragraph()
                set_spacing(p, 0, 6)
                add_runs(p, text)
        elif kind == "bullets":
            for text in items:
                bullet(doc, text)
        elif kind == "labeled":
            for it in items:
                p = doc.add_paragraph()
                set_spacing(p, 0, 3)
                add_runs(p, f"**{it['label']}:** {it['text']}")
        elif kind == "roles":
            for role in items:
                t = doc.add_paragraph()
                set_spacing(t, 6, 0)
                t.paragraph_format.keep_with_next = True
                t.paragraph_format.page_break_before = bool(role.get("break_before"))
                add_runs(t, role["title"], bold_all=True, color=BLUE)
                o = doc.add_paragraph()
                set_spacing(o, 0, 3)
                o.paragraph_format.keep_with_next = True
                add_runs(o, role["org_line"], bold_all=True, color=BLUE)
                if role.get("summary"):
                    s = doc.add_paragraph()
                    set_spacing(s, 0, 3)
                    s.paragraph_format.keep_with_next = True
                    add_runs(s, role["summary"])
                for b in role.get("bullets", []):
                    bullet(doc, b)
        else:
            raise ValueError(f"unknown section type {kind}")

    doc.save(out)
    return out


def to_pdf(docx_path):
    outdir = str(Path(docx_path).parent)
    # LibreOffice sometimes fails on its first launch or when another conversion is running; retry.
    import shutil
    import time
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("PDF conversion needs LibreOffice: 'soffice' is not on PATH")
    for attempt in range(3):
        try:
            r = subprocess.run([soffice, "--headless", "--convert-to", "pdf", str(docx_path), "--outdir", outdir],
                               capture_output=True, timeout=180)
        except subprocess.TimeoutExpired:
            time.sleep(3)
            continue
        if r.returncode == 0 and Path(docx_path).with_suffix(".pdf").exists():
            break
        time.sleep(3)
    else:
        raise RuntimeError("PDF conversion failed after 3 attempts (timed out or LibreOffice returned an error)")
    pdf = Path(docx_path).with_suffix(".pdf")
    pages = None
    try:
        info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
        pages = int(re.search(r"Pages:\s+(\d+)", info).group(1))
    except Exception:
        pass
    return pdf, pages


def main(argv=None):
    ap = argparse.ArgumentParser(description="Build a resume or cover letter .docx from a JSON spec.")
    ap.add_argument("spec", help="spec JSON (format documented at the top of this file)")
    ap.add_argument("out", help="output .docx path")
    ap.add_argument("--pdf", action="store_true", help="also convert to PDF with LibreOffice")
    ap.add_argument("--profile", help="profile.json (default: search $CAREER_AGENCY_HOME, ./, ~/.career-agency/)")
    args = ap.parse_args(argv)
    spec = json.loads(Path(args.spec).read_text(encoding="utf-8"))
    out = build(spec, args.out, find_profile(args.profile))
    print(f"wrote {out}")
    if args.pdf:
        pdf, pages = to_pdf(out)
        print(f"wrote {pdf} ({pages} pages)")


if __name__ == "__main__":
    main()
